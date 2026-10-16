from flask import (
    Blueprint,
    render_template,
    request,
    redirect,
    url_for,
    session,
    flash,
    send_file
)

from config import get_db_connection

from io import BytesIO

from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# BLUEPRINT
# ============================================================

kra_kpi_bp = Blueprint("kra_kpi", __name__)

# ============================================================
# AUDIT LOG HELPER
# ============================================================

def log_audit(user_id, action, description=None):
    conn = None
    cursor = None

    try:
        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
            INSERT INTO audit_logs
            (
                user_id,
                action,
                description,
                ip_address
            )
            VALUES
            (
                %s,
                %s,
                %s,
                %s
            )
        """, (
            user_id,
            action,
            description,
            request.remote_addr
        ))

        conn.commit()

    except Exception as e:
        print("AUDIT LOG ERROR:", e)

        if conn:
            conn.rollback()

    finally:
        if cursor:
            cursor.close()

        if conn:
            conn.close()

@kra_kpi_bp.before_request
def require_kra_kpi_administrator():

    if "user_id" not in session:
        return redirect(url_for("login_page"))

    role = session.get("role_name", "").lower()

    # Only HR and Admin can access all KRA/KPI functions (including KPI Records)
    if any(k in role for k in ["hr", "admin", "human"]):
        return None

    return "Access Denied", 403
# ============================================================
# APPRAISAL CYCLE MANAGEMENT
# ============================================================

@kra_kpi_bp.route("/appraisal_cycles.html")
@kra_kpi_bp.route("/appraisal-cycles")
def appraisal_cycles():

    conn = None
    cursor = None

    try:

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT
                cycle_id,
                cycle_name,
                start_date,
                end_date,
                status,
                created_at
            FROM appraisal_cycles
            ORDER BY start_date DESC, cycle_id DESC
        """)

        cycles = cursor.fetchall()

        return render_template(
            "appraisal_cycles.html",
            cycles=cycles
        )

    except Exception as e:

        return (
            f"Error loading Appraisal Cycles page: {str(e)}",
            500
        )

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()


# ============================================================
# ADD APPRAISAL CYCLE
# ============================================================

@kra_kpi_bp.route(
    "/add-appraisal-cycle",
    methods=["POST"]
)
def add_appraisal_cycle():

    conn = None
    cursor = None

    try:

        cycle_name = request.form.get(
            "cycle_name",
            ""
        ).strip()

        start_date = request.form.get(
            "start_date",
            ""
        ).strip()

        end_date = request.form.get(
            "end_date",
            ""
        ).strip()

        status = request.form.get(
            "status",
            "Active"
        ).strip()

        # ----------------------------------------------------
        # Validate required fields
        # ----------------------------------------------------

        if not cycle_name:
            return "Cycle name is required.", 400

        if not start_date:
            return "Start date is required.", 400

        if not end_date:
            return "End date is required.", 400

        if status not in {
            "Active",
            "Inactive",
            "Closed"
        }:
            return "Invalid appraisal cycle status.", 400

        # ----------------------------------------------------
        # Validate dates
        # ----------------------------------------------------

        if end_date < start_date:
            return (
                "End date cannot be earlier than start date.",
                400
            )

        # ----------------------------------------------------
        # Database connection
        # ----------------------------------------------------

        conn = get_db_connection()
        cursor = conn.cursor()

        # ----------------------------------------------------
        # Insert cycle
        # ----------------------------------------------------

        cursor.execute("""
            INSERT INTO appraisal_cycles
            (
                cycle_name,
                start_date,
                end_date,
                status
            )
            VALUES
            (
                %s,
                %s,
                %s,
                %s
            )
            RETURNING cycle_id
        """, (
            cycle_name,
            start_date,
            end_date,
            status
        ))

        cycle_id = cursor.fetchone()[0]

        # ----------------------------------------------------
        # Save
        # ----------------------------------------------------

        conn.commit()

        # ----------------------------------------------------
        # Audit log
        # ----------------------------------------------------

        log_audit(
            session["user_id"],
            "Appraisal Cycle Created",
            (
                f"Appraisal cycle '{cycle_name}' "
                f"(ID {cycle_id}) created"
            )
        )

        flash(
            "Appraisal cycle created successfully.",
            "success"
        )

        return redirect(
            url_for("kra_kpi.appraisal_cycles")
        )

    except Exception as e:

        if conn:
            conn.rollback()

        return (
            f"Error creating appraisal cycle: {str(e)}",
            500
        )

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()


# ============================================================
# EDIT APPRAISAL CYCLE
# ============================================================

@kra_kpi_bp.route(
    "/edit-appraisal-cycle/<int:cycle_id>",
    methods=["POST"]
)
def edit_appraisal_cycle(cycle_id):

    conn = None
    cursor = None

    try:

        cycle_name = request.form.get(
            "cycle_name",
            ""
        ).strip()

        start_date = request.form.get(
            "start_date",
            ""
        ).strip()

        end_date = request.form.get(
            "end_date",
            ""
        ).strip()

        status = request.form.get(
            "status",
            ""
        ).strip()

        # ----------------------------------------------------
        # Validate required fields
        # ----------------------------------------------------

        if not cycle_name:
            return "Cycle name is required.", 400

        if not start_date:
            return "Start date is required.", 400

        if not end_date:
            return "End date is required.", 400

        if status not in {
            "Active",
            "Inactive",
            "Closed"
        }:
            return "Invalid appraisal cycle status.", 400

        # ----------------------------------------------------
        # Validate dates
        # ----------------------------------------------------

        if end_date < start_date:
            return (
                "End date cannot be earlier than start date.",
                400
            )

        # ----------------------------------------------------
        # Database connection
        # ----------------------------------------------------

        conn = get_db_connection()
        cursor = conn.cursor()

        # ----------------------------------------------------
        # Check cycle exists
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                cycle_name
            FROM appraisal_cycles
            WHERE cycle_id = %s
        """, (cycle_id,))

        existing_cycle = cursor.fetchone()

        if not existing_cycle:
            return (
                "Appraisal cycle not found.",
                404
            )

        # ----------------------------------------------------
        # Update cycle
        # ----------------------------------------------------

        cursor.execute("""
            UPDATE appraisal_cycles
            SET
                cycle_name = %s,
                start_date = %s,
                end_date = %s,
                status = %s
            WHERE cycle_id = %s
        """, (
            cycle_name,
            start_date,
            end_date,
            status,
            cycle_id
        ))

        # ----------------------------------------------------
        # Save
        # ----------------------------------------------------

        conn.commit()

        # ----------------------------------------------------
        # Audit log
        # ----------------------------------------------------

        log_audit(
            session["user_id"],
            "Appraisal Cycle Updated",
            (
                f"Appraisal cycle ID {cycle_id} "
                f"updated to '{cycle_name}'"
            )
        )

        flash(
            "Appraisal cycle updated successfully.",
            "success"
        )

        return redirect(
            url_for("kra_kpi.appraisal_cycles")
        )

    except Exception as e:

        if conn:
            conn.rollback()

        return (
            f"Error updating appraisal cycle: {str(e)}",
            500
        )

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()
# ============================================================
# KRA MANAGEMENT
# ============================================================

@kra_kpi_bp.route("/kra_management.html")
@kra_kpi_bp.route("/kra-management")
def kra_management():

    conn = None
    cursor = None

    try:

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT
                kra_id,
                kra_name,
                description,
                default_weightage,
                created_at
            FROM kra_master
            ORDER BY kra_id DESC
        """)

        kras = cursor.fetchall()

        return render_template(
            "kra_management.html",
            kras=kras
        )

    except Exception as e:

        return f"Error loading KRA Management page: {str(e)}", 500

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()


# ============================================================
# ADD KRA
# ============================================================

@kra_kpi_bp.route("/add-kra", methods=["POST"])
def add_kra():

    conn = None
    cursor = None

    try:

        kra_name = request.form.get("kra_name")
        description = request.form.get("description")
        default_weightage = request.form.get("default_weightage")

        if not kra_name:
            return "KRA Name is required.", 400

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
            INSERT INTO kra_master
            (
                kra_name,
                description,
                default_weightage
            )
            VALUES (%s, %s, %s)
        """, (
            kra_name,
            description,
            default_weightage
        ))

        conn.commit()

        flash("KRA added successfully.", "success")

        return redirect(
            url_for("kra_kpi.kra_management")
        )

    except Exception as e:

        if conn:
            conn.rollback()

        return f"Error adding KRA: {str(e)}", 500

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()


# ============================================================
# EDIT KRA
# ============================================================

@kra_kpi_bp.route("/edit-kra/<int:kra_id>", methods=["POST"])
def edit_kra(kra_id):

    conn = None
    cursor = None

    try:

        kra_name = request.form.get("kra_name")
        description = request.form.get("description")
        default_weightage = request.form.get("default_weightage")

        if not kra_name:
            return "KRA Name is required.", 400

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
            UPDATE kra_master
            SET
                kra_name = %s,
                description = %s,
                default_weightage = %s
            WHERE kra_id = %s
        """, (
            kra_name,
            description,
            default_weightage,
            kra_id
        ))

        conn.commit()

        flash("KRA updated successfully.", "success")

        return redirect(
            url_for("kra_kpi.kra_management")
        )

    except Exception as e:

        if conn:
            conn.rollback()

        return f"Error updating KRA: {str(e)}", 500

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()


# ============================================================
# DELETE KRA
# ============================================================

@kra_kpi_bp.route("/delete-kra/<int:kra_id>", methods=["POST"])
def delete_kra(kra_id):

    conn = None
    cursor = None

    try:

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
            DELETE FROM kra_master
            WHERE kra_id = %s
        """, (kra_id,))

        conn.commit()

        flash("KRA deleted successfully.", "success")

        return redirect(
            url_for("kra_kpi.kra_management")
        )

    except Exception as e:

        if conn:
            conn.rollback()

        return f"Error deleting KRA: {str(e)}", 500

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()


# ============================================================
# ASSIGN KRA PAGE
# ============================================================

@kra_kpi_bp.route("/assign_kra.html")
@kra_kpi_bp.route("/assign-kra")
def assign_kra():

    conn = None
    cursor = None

    try:

        conn = get_db_connection()
        cursor = conn.cursor()

        # ----------------------------------------------------
        # LOAD EMPLOYEES
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                employee_id,
                first_name,
                last_name,
                employee_code
            FROM employees
            WHERE employment_status IS NULL
               OR employment_status = 'Active'
            ORDER BY first_name, last_name
        """)

        employees = cursor.fetchall()

        # ----------------------------------------------------
        # LOAD KRAs
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                kra_id,
                kra_name,
                default_weightage
            FROM kra_master
            ORDER BY kra_name
        """)

        kras = cursor.fetchall()

        # ----------------------------------------------------
        # LOAD APPRAISAL CYCLES
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                cycle_id,
                cycle_name,
                start_date,
                end_date,
                status
            FROM appraisal_cycles
            ORDER BY start_date DESC, cycle_id DESC
        """)

        cycles = cursor.fetchall()

        # ----------------------------------------------------
        # LOAD KRA ASSIGNMENTS
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                ek.employee_kra_id,
                e.first_name,
                e.last_name,
                e.employee_code,
                km.kra_name,
                ac.cycle_name,
                ek.assigned_weightage,
                ek.target_value,
                ek.due_date,
                ek.status

            FROM employee_kra ek

            INNER JOIN employees e
                ON ek.employee_id = e.employee_id

            INNER JOIN kra_master km
                ON ek.kra_id = km.kra_id

            LEFT JOIN appraisal_cycles ac
                ON ek.cycle_id = ac.cycle_id

            ORDER BY ek.employee_kra_id DESC
        """)

        assignments = cursor.fetchall()

        return render_template(
            "assign_kra.html",
            employees=employees,
            kras=kras,
            cycles=cycles,
            assignments=assignments
        )

    except Exception as e:

        return f"Error loading Assign KRA page: {str(e)}", 500

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()

# ============================================================
# SAVE ASSIGNED KRA
# ============================================================

@kra_kpi_bp.route("/assign-kra", methods=["POST"])
def save_assign_kra():

    conn = None
    cursor = None

    try:

        # ----------------------------------------------------
        # GET FORM VALUES
        # ----------------------------------------------------

        employee_id = request.form.get("employee_id")
        cycle_id = request.form.get("cycle_id")
        kra_id = request.form.get("kra_id")
        assigned_weightage = request.form.get("assigned_weightage")
        target_value = request.form.get("target_value")
        due_date = request.form.get("due_date")

        # ----------------------------------------------------
        # VALIDATE REQUIRED FIELDS
        # ----------------------------------------------------

        if not employee_id:
            return "Please select an employee.", 400

        if not cycle_id:
            return "Please select an appraisal cycle.", 400

        if not kra_id:
            return "Please select a KRA.", 400

        if not assigned_weightage:
            return "Assigned weightage is required.", 400

        if not due_date:
            return "Due date is required.", 400

        # ----------------------------------------------------
        # DATABASE CONNECTION
        # ----------------------------------------------------

        conn = get_db_connection()
        cursor = conn.cursor()

        # ----------------------------------------------------
        # VERIFY APPRAISAL CYCLE
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                cycle_id,
                start_date,
                end_date,
                status
            FROM appraisal_cycles
            WHERE cycle_id = %s
        """, (cycle_id,))

        cycle = cursor.fetchone()

        if not cycle:
            return "Selected appraisal cycle does not exist.", 400

        cycle_db_id, cycle_start, cycle_end, cycle_status = cycle

        # Do not allow assignments to inactive/closed cycles
        if cycle_status != "Active":
            return "Only an Active appraisal cycle can be assigned.", 400

        # ----------------------------------------------------
        # VERIFY DUE DATE IS WITHIN CYCLE
        # ----------------------------------------------------

        if due_date < str(cycle_start) or due_date > str(cycle_end):
            return (
                "Due date must be within the selected appraisal cycle.",
                400
            )

        # ----------------------------------------------------
        # VERIFY EMPLOYEE
        # ----------------------------------------------------

        cursor.execute("""
            SELECT employee_id
            FROM employees
            WHERE employee_id = %s
        """, (employee_id,))

        if not cursor.fetchone():
            return "Selected employee does not exist.", 400

        # ----------------------------------------------------
        # VERIFY KRA
        # ----------------------------------------------------

        cursor.execute("""
            SELECT kra_id
            FROM kra_master
            WHERE kra_id = %s
        """, (kra_id,))

        if not cursor.fetchone():
            return "Selected KRA does not exist.", 400

        # ----------------------------------------------------
        # INSERT KRA ASSIGNMENT
        # ----------------------------------------------------

        cursor.execute("""
            INSERT INTO employee_kra
            (
                employee_id,
                kra_id,
                cycle_id,
                assigned_weightage,
                target_value,
                due_date,
                status
            )
            VALUES
            (
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s
            )
        """, (
            employee_id,
            kra_id,
            cycle_id,
            assigned_weightage,
            target_value,
            due_date,
            "Active"
        ))

        # ----------------------------------------------------
        # SAVE
        # ----------------------------------------------------

        conn.commit()

        # ----------------------------------------------------
        # AUDIT LOG
        # ----------------------------------------------------

        log_audit(
            session.get("user_id"),
            "KRA Assigned",
            (
                f"KRA ID {kra_id} assigned to Employee ID "
                f"{employee_id} for Appraisal Cycle ID {cycle_id}"
            )
        )

        flash("KRA assigned successfully!", "success")

        return redirect(
            url_for("kra_kpi.assign_kra")
        )

    except Exception as e:

        if conn:
            conn.rollback()

        return f"Error assigning KRA: {str(e)}", 500

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()

# ============================================================
# UPDATE KRA STATUS
# ============================================================

@kra_kpi_bp.route(
    "/update-kra-status/<int:employee_kra_id>",
    methods=["POST"]
)
def update_kra_status(employee_kra_id):

    conn = None
    cursor = None

    try:
        new_status = request.form.get(
            "status",
            ""
        ).strip()

        # ----------------------------------------------------
        # VALIDATE STATUS
        # ----------------------------------------------------

        allowed_statuses = {
            "Active",
            "Completed",
            "Pending"
        }

        if new_status not in allowed_statuses:
            return "Invalid KRA status.", 400

        # ----------------------------------------------------
        # DATABASE CONNECTION
        # ----------------------------------------------------

        conn = get_db_connection()
        cursor = conn.cursor()

        # ----------------------------------------------------
        # CHECK KRA ASSIGNMENT EXISTS
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                employee_id,
                kra_id,
                status
            FROM employee_kra
            WHERE employee_kra_id = %s
        """, (
            employee_kra_id,
        ))

        kra_assignment = cursor.fetchone()

        if not kra_assignment:
            return "KRA assignment not found.", 404

        employee_id = kra_assignment[0]
        kra_id = kra_assignment[1]
        old_status = kra_assignment[2]

        # ----------------------------------------------------
        # UPDATE STATUS
        # ----------------------------------------------------

        cursor.execute("""
            UPDATE employee_kra
            SET status = %s
            WHERE employee_kra_id = %s
        """, (
            new_status,
            employee_kra_id
        ))

        conn.commit()

        # ----------------------------------------------------
        # AUDIT LOG
        # ----------------------------------------------------

        log_audit(
            session.get("user_id"),
            "KRA Status Updated",
            (
                f"KRA assignment ID {employee_kra_id} "
                f"for Employee ID {employee_id} "
                f"changed from '{old_status}' "
                f"to '{new_status}'"
            )
        )

        flash(
            f"KRA status updated to {new_status}.",
            "success"
        )

        return redirect(
            url_for("kra_kpi.assign_kra")
        )

    except Exception as e:

        if conn:
            conn.rollback()

        return (
            f"Error updating KRA status: {str(e)}",
            500
        )

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()

# ============================================================
# DELETE KRA ASSIGNMENT
# ============================================================

@kra_kpi_bp.route(
    "/delete-kra-assignment/<int:employee_kra_id>",
    methods=["POST"]
)
def delete_kra_assignment(employee_kra_id):

    conn = None
    cursor = None

    try:

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
            DELETE FROM employee_kra
            WHERE employee_kra_id = %s
        """, (employee_kra_id,))

        conn.commit()

        # Record the successful KRA deletion
        log_audit(
            session.get("user_id"),
            "KRA Assignment Deleted",
            f"KRA assignment ID {employee_kra_id} deleted"
        )

        return redirect(
            url_for("kra_kpi.assign_kra")
        )

    except Exception as e:

        if conn:
            conn.rollback()

        return f"Error deleting KRA assignment: {str(e)}", 500

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()


# ============================================================
# KPI MANAGEMENT
# ============================================================

@kra_kpi_bp.route("/kpi_management.html")
@kra_kpi_bp.route("/kpi-management")
def kpi_management():

    conn = None
    cursor = None

    try:

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT
                kpi_id,
                kpi_name,
                description,
                unit
            FROM kpi_master
            ORDER BY kpi_id DESC
        """)

        kpis = cursor.fetchall()

        return render_template(
            "kpi_management.html",
            kpis=kpis
        )

    except Exception as e:

        return f"Error loading KPI Management page: {str(e)}", 500

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()


# ============================================================
# ADD KPI
# ============================================================

@kra_kpi_bp.route("/add-kpi", methods=["POST"])
def add_kpi():

    conn = None
    cursor = None

    try:

        kpi_name = request.form.get("kpi_name")
        description = request.form.get("description")
        unit = request.form.get("unit")

        if not kpi_name:
            return "KPI Name is required.", 400

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
            INSERT INTO kpi_master
            (
                kpi_name,
                description,
                unit
            )
            VALUES (%s, %s, %s)
        """, (
            kpi_name,
            description,
            unit
        ))

        conn.commit()

        # Record the successful KPI creation
        log_audit(
            session.get("user_id"),
            "KPI Created",
            f"KPI '{kpi_name}' created"
        )

        return redirect(
            url_for("kra_kpi.kpi_management")
        )

    except Exception as e:

        if conn:
            conn.rollback()

        return f"Error adding KPI: {str(e)}", 500

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()


# ============================================================
# EDIT KPI
# ============================================================

@kra_kpi_bp.route("/edit-kpi/<int:kpi_id>", methods=["POST"])
def edit_kpi(kpi_id):

    conn = None
    cursor = None

    try:

        kpi_name = request.form.get("kpi_name")
        description = request.form.get("description")
        unit = request.form.get("unit")

        if not kpi_name:
            return "KPI Name is required.", 400

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
            UPDATE kpi_master
            SET
                kpi_name = %s,
                description = %s,
                unit = %s
            WHERE kpi_id = %s
        """, (
            kpi_name,
            description,
            unit,
            kpi_id
        ))

        conn.commit()

        return redirect(
            url_for("kra_kpi.kpi_management")
        )

    except Exception as e:

        if conn:
            conn.rollback()

        return f"Error updating KPI: {str(e)}", 500

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()


# ============================================================
# DELETE KPI
# ============================================================

@kra_kpi_bp.route("/delete-kpi/<int:kpi_id>", methods=["POST"])
def delete_kpi(kpi_id):

    conn = None
    cursor = None

    try:

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
            DELETE FROM kpi_master
            WHERE kpi_id = %s
        """, (kpi_id,))

        conn.commit()

        return redirect(
            url_for("kra_kpi.kpi_management")
        )

    except Exception as e:

        if conn:
            conn.rollback()

        return f"Error deleting KPI: {str(e)}", 500

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()


# ============================================================
# ASSIGN KPI PAGE
# ============================================================

@kra_kpi_bp.route("/assign_kpi.html")
@kra_kpi_bp.route("/assign-kpi")
def assign_kpi():

    conn = None
    cursor = None

    try:

        conn = get_db_connection()
        cursor = conn.cursor()

        # ----------------------------------------------------
        # LOAD EMPLOYEES
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                employee_id,
                first_name,
                last_name,
                employee_code
            FROM employees
            WHERE employment_status IS NULL
               OR employment_status = 'Active'
            ORDER BY first_name, last_name
        """)

        employees = cursor.fetchall()

        # ----------------------------------------------------
        # LOAD KPIs
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                kpi_id,
                kpi_name,
                unit
            FROM kpi_master
            ORDER BY kpi_name
        """)

        kpis = cursor.fetchall()

        # ----------------------------------------------------
        # LOAD APPRAISAL CYCLES
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                cycle_id,
                cycle_name,
                start_date,
                end_date,
                status
            FROM appraisal_cycles
            ORDER BY start_date DESC, cycle_id DESC
        """)

        cycles = cursor.fetchall()

        # ----------------------------------------------------
        # LOAD KPI ASSIGNMENTS
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                ek.employee_kpi_id,
                e.first_name,
                e.last_name,
                e.employee_code,
                km.kpi_name,
                ac.cycle_name,
                ek.assigned_target,
                ek.appraisal_period,
                ek.start_date,
                ek.end_date,
                ek.status

            FROM employee_kpi ek

            INNER JOIN employees e
                ON ek.employee_id = e.employee_id

            INNER JOIN kpi_master km
                ON ek.kpi_id = km.kpi_id

            LEFT JOIN appraisal_cycles ac
                ON ek.cycle_id = ac.cycle_id

            ORDER BY ek.employee_kpi_id DESC
        """)

        assignments = cursor.fetchall()

        return render_template(
            "assign_kpi.html",
            employees=employees,
            kpis=kpis,
            cycles=cycles,
            assignments=assignments
        )

    except Exception as e:

        return f"Error loading Assign KPI page: {str(e)}", 500

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()
# ============================================================
# SAVE ASSIGNED KPI
# ============================================================

@kra_kpi_bp.route("/assign-kpi", methods=["POST"])
def save_assign_kpi():

    conn = None
    cursor = None

    try:

        # ----------------------------------------------------
        # GET FORM VALUES
        # ----------------------------------------------------

        employee_id = request.form.get("employee_id")
        cycle_id = request.form.get("cycle_id")
        kpi_id = request.form.get("kpi_id")
        assigned_target = request.form.get("assigned_target")
        appraisal_period = request.form.get("appraisal_period")
        start_date = request.form.get("start_date")
        end_date = request.form.get("end_date")

        # ----------------------------------------------------
        # VALIDATE REQUIRED FIELDS
        # ----------------------------------------------------

        if not employee_id:
            return "Please select an employee.", 400

        if not cycle_id:
            return "Please select an appraisal cycle.", 400

        if not kpi_id:
            return "Please select a KPI.", 400

        if not assigned_target:
            return "Target value is required.", 400

        if not appraisal_period:
            return "Appraisal period is required.", 400

        if not start_date:
            return "Start date is required.", 400

        if not end_date:
            return "End date is required.", 400

        # ----------------------------------------------------
        # VALIDATE DATE ORDER
        # ----------------------------------------------------

        if end_date < start_date:
            return "End date cannot be earlier than start date.", 400

        # ----------------------------------------------------
        # DATABASE CONNECTION
        # ----------------------------------------------------

        conn = get_db_connection()
        cursor = conn.cursor()

        # ----------------------------------------------------
        # VERIFY APPRAISAL CYCLE
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                cycle_id,
                start_date,
                end_date,
                status
            FROM appraisal_cycles
            WHERE cycle_id = %s
        """, (cycle_id,))

        cycle = cursor.fetchone()

        if not cycle:
            return "Selected appraisal cycle does not exist.", 400

        cycle_db_id, cycle_start, cycle_end, cycle_status = cycle

        # Only Active cycles can receive new assignments
        if cycle_status != "Active":
            return "Only an Active appraisal cycle can be assigned.", 400

        # ----------------------------------------------------
        # VERIFY KPI DATES ARE WITHIN THE CYCLE
        # ----------------------------------------------------

        if start_date < str(cycle_start) or start_date > str(cycle_end):
            return (
                "Start date must be within the selected appraisal cycle.",
                400
            )

        if end_date < str(cycle_start) or end_date > str(cycle_end):
            return (
                "End date must be within the selected appraisal cycle.",
                400
            )

        # ----------------------------------------------------
        # VERIFY EMPLOYEE
        # ----------------------------------------------------

        cursor.execute("""
            SELECT employee_id
            FROM employees
            WHERE employee_id = %s
        """, (employee_id,))

        if not cursor.fetchone():
            return "Selected employee does not exist.", 400

        # ----------------------------------------------------
        # VERIFY KPI
        # ----------------------------------------------------

        cursor.execute("""
            SELECT kpi_id
            FROM kpi_master
            WHERE kpi_id = %s
        """, (kpi_id,))

        if not cursor.fetchone():
            return "Selected KPI does not exist.", 400

        # ----------------------------------------------------
        # GET LOGGED-IN USER
        # ----------------------------------------------------

        assigned_by = session.get("user_id")

        # ----------------------------------------------------
        # INSERT KPI ASSIGNMENT
        # ----------------------------------------------------

        cursor.execute("""
            INSERT INTO employee_kpi
            (
                employee_id,
                kpi_id,
                cycle_id,
                assigned_target,
                appraisal_period,
                start_date,
                end_date,
                status,
                assigned_by,
                assigned_at
            )
            VALUES
            (
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                CURRENT_TIMESTAMP
            )
        """, (
            employee_id,
            kpi_id,
            cycle_id,
            assigned_target,
            appraisal_period,
            start_date,
            end_date,
            "Active",
            assigned_by
        ))

        # ----------------------------------------------------
        # SAVE
        # ----------------------------------------------------

        conn.commit()

        # ----------------------------------------------------
        # AUDIT LOG
        # ----------------------------------------------------

        log_audit(
            session.get("user_id"),
            "KPI Assigned",
            (
                f"KPI ID {kpi_id} assigned to Employee ID "
                f"{employee_id} for Appraisal Cycle ID {cycle_id}"
            )
        )

        flash("KPI assigned successfully!", "success")

        return redirect(
            url_for("kra_kpi.assign_kpi")
        )

    except Exception as e:

        if conn:
            conn.rollback()

        print("Error assigning KPI:", e)

        return (
            f"Error assigning KPI: {str(e)}",
            500
        )

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()
# ============================================================
# DELETE KPI ASSIGNMENT
# ============================================================

@kra_kpi_bp.route(
    "/delete-kpi-assignment/<int:employee_kpi_id>",
    methods=["POST"]
)
def delete_kpi_assignment(employee_kpi_id):

    conn = None
    cursor = None

    try:

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
            DELETE FROM employee_kpi
            WHERE employee_kpi_id = %s
        """, (employee_kpi_id,))

        conn.commit()

        # Record the successful KPI assignment deletion
        log_audit(
            session.get("user_id"),
            "KPI Assignment Deleted",
            f"KPI assignment ID {employee_kpi_id} deleted"
        )

        return redirect(
            url_for("kra_kpi.assign_kpi")
        )
    except Exception as e:

        if conn:
            conn.rollback()

        return f"Error deleting KPI assignment: {str(e)}", 500

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()


# ============================================================
# KPI RECORDS PAGE
# ============================================================

@kra_kpi_bp.route("/kpi_records.html")
@kra_kpi_bp.route("/kpi-records")
def kpi_records():

    conn = None
    cursor = None

    try:

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT
                kr.employee_kpi_id,

                e.employee_code,
                e.first_name,
                e.last_name,

                km.kpi_name,

                ac.cycle_name,

                kr.achieved_value,
                kr.evaluation_month,
                kr.remarks

            FROM kpi_records kr

            INNER JOIN employees e
                ON kr.employee_id = e.employee_id

            INNER JOIN kpi_master km
                ON kr.kpi_id = km.kpi_id

            LEFT JOIN employee_kpi ek
                ON kr.employee_id = ek.employee_id
                AND kr.kpi_id = ek.kpi_id
                AND kr.evaluation_month BETWEEN ek.start_date AND ek.end_date

            LEFT JOIN appraisal_cycles ac
                ON ek.cycle_id = ac.cycle_id

            ORDER BY
                kr.evaluation_month DESC,
                kr.employee_kpi_id DESC
        """)
        records = cursor.fetchall()

        cursor.execute("""
            SELECT
                employee_id,
                first_name,
                last_name,
                employee_code
            FROM employees
            WHERE employment_status IS NULL
               OR employment_status = 'Active'
            ORDER BY first_name, last_name
        """)

        employees = cursor.fetchall()

        cursor.execute("""
            SELECT
                kpi_id,
                kpi_name,
                unit
            FROM kpi_master
            ORDER BY kpi_name
        """)

        kpis = cursor.fetchall()

        return render_template(
            "kpi_records.html",
            records=records,
            employees=employees,
            kpis=kpis
        )

    except Exception as e:

        return f"Error loading KPI Records page: {str(e)}", 500

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()


# ============================================================
# ADD KPI RECORD
# ============================================================

@kra_kpi_bp.route("/add-kpi-record", methods=["POST"])
def add_kpi_record():

    conn = None
    cursor = None
    from routes.dashboard import save_kpi_performance_result
    print("DEBUG: add_kpi_record reached")
    try:

        employee_id = request.form.get("employee_id")
        kpi_id = request.form.get("kpi_id")
        achieved_value = request.form.get("achieved_value")
        evaluation_month = request.form.get("evaluation_month")

        if evaluation_month:
            evaluation_month = evaluation_month + "-01"
        remarks = request.form.get("remarks")

        if not employee_id:
            return "Please select an employee.", 400

        if not kpi_id:
            return "Please select a KPI.", 400

        if not achieved_value:
            return "Achieved value is required.", 400

        if not evaluation_month:
            return "Evaluation month is required.", 400

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
            INSERT INTO kpi_records
            (
                employee_id,
                kpi_id,
                achieved_value,
                evaluation_month,
                remarks
            )
            VALUES
            (
                %s,
                %s,
                %s,
                %s,
                %s
            )
        """, (
            employee_id,
            kpi_id,
            achieved_value,
            evaluation_month,
            remarks
        ))

        # ----------------------------------------------------
        # FIND THE APPRAISAL CYCLE FOR THIS KPI RECORD
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                cycle_id
            FROM employee_kpi
            WHERE employee_id = %s
              AND kpi_id = %s
              AND DATE_TRUNC('month', start_date) <= DATE_TRUNC('month', %s::date)
              AND DATE_TRUNC('month', end_date) >= DATE_TRUNC('month', %s::date)
            ORDER BY employee_kpi_id DESC
            LIMIT 1
        """, (
            employee_id,
            kpi_id,
            evaluation_month,
            evaluation_month
        ))

        assignment = cursor.fetchone()

        # ----------------------------------------------------
        # SAVE KPI SCORE TO PERFORMANCE RESULT
        # ----------------------------------------------------

        if assignment and assignment[0]:
            cycle_id = assignment[0]

            save_kpi_performance_result(
                cursor,
                int(employee_id),
                int(cycle_id)
            )

        # ----------------------------------------------------
        # SAVE EVERYTHING
        # ----------------------------------------------------

        conn.commit()

        # Record the successful KPI record creation
        log_audit(
            session.get("user_id"),
            "KPI Record Added",
            f"KPI record added for Employee ID {employee_id}, KPI ID {kpi_id}"
        )

        flash("KPI record added successfully!", "success")

        return redirect(
            url_for("kra_kpi.kpi_records")
        )

    except Exception as e:

        if conn:
            conn.rollback()

        return f"Error adding KPI record: {str(e)}", 500

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()


# ============================================================
# DELETE KPI RECORD
# ============================================================

@kra_kpi_bp.route(
    "/delete-kpi-record/<int:employee_kpi_id>",
    methods=["POST"]
)
def delete_kpi_record(employee_kpi_id):

    conn = None
    cursor = None

    try:

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
            DELETE FROM kpi_records
            WHERE employee_kpi_id = %s
        """, (employee_kpi_id,))

        conn.commit()

        return redirect(
            url_for("kra_kpi.kpi_records")
        )

    except Exception as e:

        if conn:
            conn.rollback()

        return f"Error deleting KPI record: {str(e)}", 500

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()
# ============================================================
# KRA / KPI REPORT
# ============================================================

@kra_kpi_bp.route("/kra_kpi_report.html")
@kra_kpi_bp.route("/kra-kpi-report")
def kra_kpi_report():

    conn = None
    cursor = None

    try:
        conn = get_db_connection()
        cursor = conn.cursor()

        # ----------------------------------------------------
        # GET FILTER VALUES FROM URL
        # ----------------------------------------------------

        search = request.args.get("search", "").strip()
        department = request.args.get("department", "").strip()
        status = request.args.get("status", "").strip()
        cycle_id = request.args.get("cycle_id", "").strip()

        # ----------------------------------------------------
        # LOAD DEPARTMENTS FOR DROPDOWN
        # ----------------------------------------------------

        cursor.execute("""
            SELECT department_name
            FROM departments
            ORDER BY department_name
        """)

        departments = cursor.fetchall()
        # ----------------------------------------------------
        # LOAD APPRAISAL CYCLES FOR DROPDOWN
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                cycle_id,
                cycle_name,
                start_date,
                end_date,
                status
            FROM appraisal_cycles
            ORDER BY start_date DESC, cycle_id DESC
        """)

        cycles = cursor.fetchall()
        # ----------------------------------------------------
        # KRA REPORT
        # ----------------------------------------------------

        kra_query = """
            SELECT
                e.employee_code,
                e.first_name,
                e.last_name,
                km.kra_name,
                ac.cycle_name,
                ek.assigned_weightage,
                ek.target_value,
                ek.due_date,
                ek.status

            FROM employee_kra ek

            INNER JOIN employees e
                ON ek.employee_id = e.employee_id

            INNER JOIN kra_master km
                ON ek.kra_id = km.kra_id

            LEFT JOIN appraisal_cycles ac
                ON ek.cycle_id = ac.cycle_id

            LEFT JOIN departments d
                ON e.department_id = d.department_id

            WHERE 1 = 1
        """

        kra_params = []

        # Employee name / ID search
        if search:
            kra_query += """
                AND (
                    LOWER(e.first_name || ' ' || e.last_name)
                    LIKE LOWER(%s)
                    OR LOWER(e.employee_code)
                    LIKE LOWER(%s)
                )
            """

            search_value = f"%{search}%"

            kra_params.extend([
                search_value,
                search_value
            ])

        # Department filter
        if department:
            kra_query += """
                AND d.department_name = %s
            """

            kra_params.append(department)

        # Status filter
        if status:
            kra_query += """
                AND ek.status = %s
            """

            kra_params.append(status)

        # Appraisal Cycle filter
        if cycle_id:
            kra_query += """
                AND ek.cycle_id = %s
            """

            kra_params.append(cycle_id)

        kra_query += """
            ORDER BY
                e.first_name,
                e.last_name,
                km.kra_name
        """

        cursor.execute(kra_query, tuple(kra_params))

        kra_report = cursor.fetchall()

        # ----------------------------------------------------
        # KPI REPORT
        # ----------------------------------------------------

        kpi_query = """
            SELECT
                e.employee_code,
                e.first_name,
                e.last_name,
                km.kpi_name,
                ac.cycle_name,
                kr.achieved_value,
                kr.evaluation_month,
                kr.remarks

            FROM kpi_records kr

            INNER JOIN employee_kpi ek
                ON kr.employee_kpi_id = ek.employee_kpi_id

            INNER JOIN employees e
                ON kr.employee_id = e.employee_id

            INNER JOIN kpi_master km
                ON kr.kpi_id = km.kpi_id

            LEFT JOIN appraisal_cycles ac
                ON ek.cycle_id = ac.cycle_id

            LEFT JOIN departments d
                ON e.department_id = d.department_id

            WHERE 1 = 1
        """

        kpi_params = []

        # Employee name / ID search
        if search:
            kpi_query += """
                AND (
                    LOWER(e.first_name || ' ' || e.last_name)
                    LIKE LOWER(%s)
                    OR LOWER(e.employee_code)
                    LIKE LOWER(%s)
                )
            """

            search_value = f"%{search}%"

            kpi_params.extend([
                search_value,
                search_value
            ])

        # Department filter
        if department:
            kpi_query += """
                AND d.department_name = %s
            """

            kpi_params.append(department)

        # Appraisal Cycle filter
        if cycle_id:
            kpi_query += """
                AND ek.cycle_id = %s
            """

            kpi_params.append(cycle_id)

        kpi_query += """
            ORDER BY
                kr.evaluation_month DESC,
                e.first_name,
                e.last_name
        """

        cursor.execute(kpi_query, tuple(kpi_params))

        kpi_report = cursor.fetchall()

        # ----------------------------------------------------
        # SEND DATA TO TEMPLATE
        # ----------------------------------------------------

        return render_template(
            "kra_kpi_report.html",
            kra_report=kra_report,
            kpi_report=kpi_report,
            departments=departments,
            cycles=cycles,
            selected_search=search,
            selected_department=department,
            selected_status=status,
            selected_cycle=cycle_id
        )

    except Exception as e:

        print("KRA/KPI REPORT ERROR:", e)

        return (
            f"Error loading KRA/KPI Report: {str(e)}",
            500
        )

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()

# ============================================================
# KRA / KPI REPORT - PDF EXPORT
# ============================================================

@kra_kpi_bp.route("/kra-kpi-report/export/pdf")
def export_kra_kpi_pdf():

    conn = None
    cursor = None

    try:

        # ----------------------------------------------------
        # GET FILTER VALUES
        # ----------------------------------------------------

        search = request.args.get("search", "").strip()
        department = request.args.get("department", "").strip()
        status = request.args.get("status", "").strip()
        cycle_id = request.args.get("cycle_id", "").strip()

        conn = get_db_connection()
        cursor = conn.cursor()

        # ----------------------------------------------------
        # KRA DATA
        # ----------------------------------------------------

        kra_query = """
            SELECT
                e.employee_code,
                e.first_name,
                e.last_name,
                km.kra_name,
                ac.cycle_name,
                ek.assigned_weightage,
                ek.target_value,
                ek.due_date,
                ek.status
            FROM employee_kra ek

            INNER JOIN employees e ON ek.employee_id = e.employee_id
            INNER JOIN kra_master km ON ek.kra_id = km.kra_id
            LEFT JOIN departments d ON e.department_id = d.department_id
            LEFT JOIN appraisal_cycles ac ON ek.cycle_id = ac.cycle_id

            WHERE 1 = 1
        """

        kra_params = []

        if search:
            kra_query += """
                AND (
                    LOWER(e.first_name || ' ' || e.last_name)
                    LIKE LOWER(%s)

                    OR LOWER(e.employee_code)
                    LIKE LOWER(%s)
                )
            """

            search_value = f"%{search}%"

            kra_params.extend([
                search_value,
                search_value
            ])

        if department:
            kra_query += """
                AND d.department_name = %s
            """

            kra_params.append(department)

        if status:
            kra_query += """
                AND ek.status = %s
            """

            kra_params.append(status)

        if cycle_id:
            kra_query += """
                AND ek.cycle_id = %s
            """
            kra_params.append(cycle_id)

        kra_query += """
            ORDER BY
                e.first_name,
                e.last_name,
                km.kra_name
        """

        cursor.execute(
            kra_query,
            tuple(kra_params)
        )

        kra_report = cursor.fetchall()

        # ----------------------------------------------------
        # KPI DATA
        # ----------------------------------------------------

        kpi_query = """
            SELECT
                e.employee_code,
                e.first_name,
                e.last_name,
                km.kpi_name,
                ac.cycle_name,
                kr.achieved_value,
                kr.evaluation_month,
                kr.remarks
            FROM kpi_records kr

            INNER JOIN employees e ON kr.employee_id = e.employee_id
            INNER JOIN kpi_master km ON kr.kpi_id = km.kpi_id
            INNER JOIN employee_kpi ek ON kr.employee_kpi_id = ek.employee_kpi_id
            LEFT JOIN appraisal_cycles ac ON ek.cycle_id = ac.cycle_id
            LEFT JOIN departments d ON e.department_id = d.department_id

            WHERE 1 = 1
        """

        kpi_params = []

        if search:
            kpi_query += """
                AND (
                    LOWER(e.first_name || ' ' || e.last_name)
                    LIKE LOWER(%s)

                    OR LOWER(e.employee_code)
                    LIKE LOWER(%s)
                )
            """

            search_value = f"%{search}%"

            kpi_params.extend([
                search_value,
                search_value
            ])

        if department:
            kpi_query += """
                AND d.department_name = %s
            """

            kpi_params.append(department)

        if cycle_id:
            kpi_query += """
                AND ek.cycle_id = %s
            """
            kpi_params.append(cycle_id)

        kpi_query += """
            ORDER BY
                kr.evaluation_month DESC,
                e.first_name,
                e.last_name
        """

        cursor.execute(
            kpi_query,
            tuple(kpi_params)
        )

        kpi_report = cursor.fetchall()

        # ----------------------------------------------------
        # CREATE PDF IN MEMORY
        # ----------------------------------------------------

        pdf_buffer = BytesIO()

        document = SimpleDocTemplate(
            pdf_buffer,
            pagesize=landscape(A4),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )

        styles = getSampleStyleSheet()

        title_style = styles["Title"]
        title_style.alignment = TA_CENTER
        title_style.fontSize = 18

        normal_style = styles["Normal"]

        elements = []

        # ----------------------------------------------------
        # TITLE
        # ----------------------------------------------------

        elements.append(
            Paragraph(
                "MUTHOOT FINANCE",
                title_style
            )
        )

        elements.append(
            Paragraph(
                "KRA / KPI REPORT",
                styles["Heading2"]
            )
        )

        elements.append(Spacer(1, 10))

        # ----------------------------------------------------
        # FILTER INFORMATION
        # ----------------------------------------------------

        filter_text = (
            f"<b>Employee Search:</b> "
            f"{search or 'All'} &nbsp;&nbsp;&nbsp; "
            f"<b>Department:</b> "
            f"{department or 'All Departments'} "
            f"&nbsp;&nbsp;&nbsp; "
            f"<b>Status:</b> "
            f"{status or 'All Statuses'} "
            f"&nbsp;&nbsp;&nbsp; "
            f"<b>Appraisal Cycle:</b> "
            f"{cycle_id or 'All Cycles'}"
        )

        elements.append(
            Paragraph(
                filter_text,
                normal_style
            )
        )

        elements.append(Spacer(1, 15))

        # ----------------------------------------------------
        # KRA SECTION
        # ----------------------------------------------------

        elements.append(
            Paragraph(
                "KRA Assignment Report",
                styles["Heading2"]
            )
        )

        kra_data = [[
            "Employee",
            "Employee ID",
            "KRA",
            "Appraisal Cycle",
            "Weightage",
            "Target",
            "Due Date",
            "Status"
        ]]

        for row in kra_report:
            kra_data.append([
                f"{row[1]} {row[2]}",
                row[0],
                row[3],
                row[4],
                row[5],
                row[6],
                row[7],
                row[8]
            ])

        if len(kra_data) == 1:
            kra_data.append([
                "No KRA records found.",
                "",
                "",
                "",
                "",
                "",
                ""
            ])

        kra_table = Table(
            kra_data,
            repeatRows=1
        )

        kra_table.setStyle(
            TableStyle([
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#C8102E")
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),
                (
                    "ALIGN",
                    (0, 0),
                    (-1, 0),
                    "CENTER"
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE"
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor("#F7F7F7")
                    ]
                )
            ])
        )

        elements.append(kra_table)

        elements.append(Spacer(1, 25))

        # ----------------------------------------------------
        # KPI SECTION
        # ----------------------------------------------------

        elements.append(
            Paragraph(
                "KPI Achievement Report",
                styles["Heading2"]
            )
        )

        kpi_data = [
            [
                "Employee",
                "Employee ID",
                "KPI",
                "Achieved Value",
                "Evaluation Month",
                "Remarks"
            ]
        ]

        for record in kpi_report:
            kpi_data.append([
                f"{record[1]} {record[2]}",
                record[0],
                record[3],
                str(record[4]),
                str(record[5]),
                str(record[6] or "-")
            ])

        if len(kpi_data) == 1:
            kpi_data.append([
                "No KPI records found.",
                "",
                "",
                "",
                "",
                ""
            ])

        kpi_table = Table(
            kpi_data,
            repeatRows=1
        )

        kpi_table.setStyle(
            TableStyle([
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#C8102E")
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),
                (
                    "ALIGN",
                    (0, 0),
                    (-1, 0),
                    "CENTER"
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE"
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor("#F7F7F7")
                    ]
                )
            ])
        )

        elements.append(kpi_table)

        # ----------------------------------------------------
        # BUILD PDF
        # ----------------------------------------------------

        document.build(elements)

        pdf_buffer.seek(0)

        return send_file(
            pdf_buffer,
            as_attachment=True,
            download_name="KRA_KPI_Report.pdf",
            mimetype="application/pdf"
        )

    except Exception as e:

        print(
            "KRA/KPI PDF EXPORT ERROR:",
            e
        )

        return (
            f"Error generating PDF report: {str(e)}",
            500
        )

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()
# ============================================================
# KRA / KPI REPORT - EXCEL EXPORT
# ============================================================

@kra_kpi_bp.route("/kra-kpi-report/export/excel")
def export_kra_kpi_excel():

    conn = None
    cursor = None

    try:

        # ----------------------------------------------------
        # GET FILTER VALUES
        # ----------------------------------------------------

        search = request.args.get("search", "").strip()
        department = request.args.get("department", "").strip()
        status = request.args.get("status", "").strip()
        cycle_id = request.args.get("cycle_id", "").strip()

        conn = get_db_connection()
        cursor = conn.cursor()

        # ----------------------------------------------------
        # KRA DATA
        # ----------------------------------------------------

        kra_query = """
            SELECT
                e.employee_code,
                e.first_name,
                e.last_name,
                km.kra_name,
                ac.cycle_name,
                ek.assigned_weightage,
                ek.target_value,
                ek.due_date,
                ek.status
            FROM employee_kra ek

            INNER JOIN employees e
                ON ek.employee_id = e.employee_id

            INNER JOIN kra_master km
                ON ek.kra_id = km.kra_id
            
            LEFT JOIN appraisal_cycles ac
                ON ek.cycle_id = ac.cycle_id

            LEFT JOIN departments d
                ON e.department_id = d.department_id

            WHERE 1 = 1
        """

        kra_params = []

        if search:

            kra_query += """
                AND (
                    LOWER(e.first_name || ' ' || e.last_name)
                    LIKE LOWER(%s)

                    OR LOWER(e.employee_code)
                    LIKE LOWER(%s)
                )
            """

            search_value = f"%{search}%"

            kra_params.extend([
                search_value,
                search_value
            ])

        if department:

            kra_query += """
                AND d.department_name = %s
            """

            kra_params.append(department)

        if status:

            kra_query += """
                AND ek.status = %s
            """

            kra_params.append(status)

        if cycle_id:
            kra_query += """
                AND ek.cycle_id = %s
            """

            kra_params.append(cycle_id)

        kra_query += """
            ORDER BY
                e.first_name,
                e.last_name,
                km.kra_name
        """

        cursor.execute(
            kra_query,
            tuple(kra_params)
        )

        kra_report = cursor.fetchall()

        # ----------------------------------------------------
        # KPI DATA
        # ----------------------------------------------------

        kpi_query = """
            SELECT
                e.employee_code,
                e.first_name,
                e.last_name,
                km.kpi_name,
                ac.cycle_name,
                kr.achieved_value,
                kr.evaluation_month,
                kr.remarks
            FROM kpi_records kr

            INNER JOIN employees e
                ON kr.employee_id = e.employee_id

            INNER JOIN kpi_master km
                ON kr.kpi_id = km.kpi_id
            
            INNER JOIN employee_kpi ek
                ON kr.employee_kpi_id = ek.employee_kpi_id
            
            LEFT JOIN appraisal_cycles ac
                ON ek.cycle_id = ac.cycle_id

            LEFT JOIN departments d
                ON e.department_id = d.department_id

            WHERE 1 = 1
        """

        kpi_params = []

        if search:

            kpi_query += """
                AND (
                    LOWER(e.first_name || ' ' || e.last_name)
                    LIKE LOWER(%s)

                    OR LOWER(e.employee_code)
                    LIKE LOWER(%s)
                )
            """

            search_value = f"%{search}%"

            kpi_params.extend([
                search_value,
                search_value
            ])

        if department:

            kpi_query += """
                AND d.department_name = %s
            """

            kpi_params.append(department)

        if cycle_id:
            kpi_query += """
                AND ek.cycle_id = %s
            """

            kpi_params.append(cycle_id)

        kpi_query += """
            ORDER BY
                kr.evaluation_month DESC,
                e.first_name,
                e.last_name
        """

        cursor.execute(
            kpi_query,
            tuple(kpi_params)
        )

        kpi_report = cursor.fetchall()

        # ----------------------------------------------------
        # CREATE WORKBOOK
        # ----------------------------------------------------

        workbook = Workbook()

        kra_sheet = workbook.active
        kra_sheet.title = "KRA Report"

        kpi_sheet = workbook.create_sheet(
            "KPI Report"
        )

        # ----------------------------------------------------
        # COMMON STYLES
        # ----------------------------------------------------

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="C8102E"
        )

        header_font = Font(
            bold=True,
            color="FFFFFF"
        )

        title_font = Font(
            bold=True,
            size=16
        )

        # ----------------------------------------------------
        # KRA SHEET
        # ----------------------------------------------------

        kra_sheet["A1"] = "MUTHOOT FINANCE"
        kra_sheet["A1"].font = title_font

        kra_sheet["A2"] = "KRA Assignment Report"
        kra_sheet["A2"].font = Font(
            bold=True,
            size=13
        )

        kra_sheet["A3"] = (
            "Employee Search"
        )

        kra_sheet["B3"] = (
            search or "All"
        )

        kra_sheet["A4"] = (
            "Department"
        )

        kra_sheet["B4"] = (
            department or "All Departments"
        )

        kra_sheet["A5"] = (
            "Status"
        )

        kra_sheet["B5"] = (
            status or "All Statuses"
        )

        kra_sheet["A6"] = "Appraisal Cycle"

        kra_sheet["B6"] = (
                cycle_id or "All Cycles"
        )

        kra_headers = [
            "Employee",
            "Employee ID",
            "KRA",
            "Appraisal Cycle",
            "Weightage",
            "Target",
            "Due Date",
            "Status"
        ]

        header_row = 7

        for column, header in enumerate(
            kra_headers,
            start=1
        ):

            cell = kra_sheet.cell(
                row=header_row,
                column=column
            )

            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center"
            )

        for row_number, record in enumerate(
            kra_report,
            start=header_row + 1
        ):

            values = [
                f"{record[1]} {record[2]}",
                record[0],
                record[3],
                record[4] or "-",
                f"{record[5]}%",
                record[6],
                record[7],
                record[8]
            ]

            for column, value in enumerate(
                values,
                start=1
            ):

                kra_sheet.cell(
                    row=row_number,
                    column=column,
                    value=value
                )

        if not kra_report:

            kra_sheet.cell(
                row=header_row + 1,
                column=1,
                value="No KRA records found."
            )

        # ----------------------------------------------------
        # KPI SHEET
        # ----------------------------------------------------

        kpi_sheet["A1"] = "MUTHOOT FINANCE"
        kpi_sheet["A1"].font = title_font

        kpi_sheet["A2"] = "KPI Achievement Report"
        kpi_sheet["A2"].font = Font(
            bold=True,
            size=13
        )

        kpi_sheet["A3"] = (
            "Employee Search"
        )

        kpi_sheet["B3"] = (
            search or "All"
        )

        kpi_sheet["A4"] = (
            "Department"
        )

        kpi_sheet["B4"] = (
            department or "All Departments"
        )

        kpi_sheet["A5"] = (
            "Status Filter"
        )

        kpi_sheet["B5"] = (
            status or "All Statuses"
        )

        kpi_sheet["A6"] = "Appraisal Cycle"

        kpi_sheet["B6"] = (
                cycle_id or "All Cycles"
        )

        kpi_headers = [
            "Employee",
            "Employee ID",
            "KPI",
            "Appraisal Cycle",
            "Achieved Value",
            "Evaluation Month",
            "Remarks"
        ]

        for column, header in enumerate(
            kpi_headers,
            start=1
        ):

            cell = kpi_sheet.cell(
                row=header_row,
                column=column
            )

            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center"
            )

        for row_number, record in enumerate(
            kpi_report,
            start=header_row + 1
        ):

            values = [
                f"{record[1]} {record[2]}",
                record[0],
                record[3],
                record[4] or "-",
                record[5],
                record[6],
                record[7] or "-"
            ]

            for column, value in enumerate(
                values,
                start=1
            ):

                kpi_sheet.cell(
                    row=row_number,
                    column=column,
                    value=value
                )

        if not kpi_report:

            kpi_sheet.cell(
                row=header_row + 1,
                column=1,
                value="No KPI records found."
            )

        # ----------------------------------------------------
        # AUTO WIDTH
        # ----------------------------------------------------

        for sheet in [
            kra_sheet,
            kpi_sheet
        ]:

            for column_cells in sheet.columns:

                max_length = 0

                column_letter = get_column_letter(
                    column_cells[0].column
                )

                for cell in column_cells:

                    if cell.value is not None:

                        max_length = max(
                            max_length,
                            len(str(cell.value))
                        )

                sheet.column_dimensions[
                    column_letter
                ].width = min(
                    max_length + 3,
                    40
                )

        # ----------------------------------------------------
        # CREATE EXCEL FILE IN MEMORY
        # ----------------------------------------------------

        excel_buffer = BytesIO()

        workbook.save(
            excel_buffer
        )

        excel_buffer.seek(0)

        return send_file(
            excel_buffer,
            as_attachment=True,
            download_name="KRA_KPI_Report.xlsx",
            mimetype=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            )
        )

    except Exception as e:

        print(
            "KRA/KPI EXCEL EXPORT ERROR:",
            e
        )

        return (
            f"Error generating Excel report: {str(e)}",
            500
        )

    finally:

        if cursor:
            cursor.close()

        if conn:
            conn.close()