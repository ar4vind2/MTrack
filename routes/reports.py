import os
from datetime import datetime
from flask import Blueprint, jsonify, request, session, send_file, current_app
from werkzeug.utils import secure_filename
from config import get_db_connection

# ReportLab Imports
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

# OpenPyXL Imports
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

reports_bp = Blueprint("reports", __name__)

REPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "static", "reports")
os.makedirs(REPORTS_DIR, exist_ok=True)


def _get_appraisal_report_data(user_id=None, role_name=None, employee_id=None, cycle_id=None):
    """Fetch structured performance appraisal data for reporting."""
    conn = get_db_connection()
    cur = conn.cursor()
    try:
        # Determine cycle_id if not explicitly provided
        selected_cycle_id = cycle_id
        selected_cycle_name = "All Cycles"
        if selected_cycle_id:
            cur.execute("SELECT cycle_name FROM appraisal_cycles WHERE cycle_id = %s", (selected_cycle_id,))
            c_row = cur.fetchone()
            if c_row:
                selected_cycle_name = c_row[0]
        else:
            cur.execute("SELECT cycle_id, cycle_name FROM appraisal_cycles WHERE status = 'Active' ORDER BY cycle_id DESC LIMIT 1")
            act_row = cur.fetchone()
            if act_row:
                selected_cycle_id, selected_cycle_name = act_row[0], act_row[1]
            else:
                cur.execute("SELECT cycle_id, cycle_name FROM appraisal_cycles ORDER BY cycle_id DESC LIMIT 1")
                max_row = cur.fetchone()
                if max_row:
                    selected_cycle_id, selected_cycle_name = max_row[0], max_row[1]

        where_clauses = []
        params = []
        r = (role_name or "").lower().strip()
        is_hr = "human" in r or "hr" in r or "admin" in r
        is_manager = "manager" in r or "lead" in r or "supervisor" in r
        is_employee = not is_hr and not is_manager

        target_emp_id = employee_id
        if not target_emp_id and user_id:
            cur.execute("SELECT employee_id FROM employees WHERE user_id = %s", (user_id,))
            emp_row = cur.fetchone()
            if emp_row:
                target_emp_id = emp_row[0]

        if is_employee:
            if target_emp_id:
                where_clauses.append("e.employee_id = %s")
                params.append(target_emp_id)
            elif user_id:
                where_clauses.append("e.user_id = %s")
                params.append(user_id)
        elif is_manager:
            if target_emp_id:
                where_clauses.append("(e.employee_id = %s OR e.manager_id = %s)")
                params.extend([target_emp_id, target_emp_id])
            elif user_id:
                where_clauses.append("(e.user_id = %s OR e.manager_id = (SELECT employee_id FROM employees WHERE user_id = %s))")
                params.extend([user_id, user_id])

        where_sql = ("WHERE " + " AND ".join(where_clauses)) if where_clauses else ""

        if selected_cycle_id:
            join_pr_sql = """
                LEFT JOIN (
                    SELECT employee_id, final_score, rating, recommendation
                    FROM performance_result
                    WHERE cycle_id = %s
                ) pr ON e.employee_id = pr.employee_id
            """
            join_params = [selected_cycle_id]
        else:
            join_pr_sql = """
                LEFT JOIN (
                    SELECT DISTINCT ON (employee_id) employee_id, final_score, rating, recommendation
                    FROM performance_result
                    ORDER BY employee_id, performance_id DESC
                ) pr ON e.employee_id = pr.employee_id
            """
            join_params = []

        query = f"""
            SELECT 
                e.employee_code,
                CONCAT(e.first_name, ' ', e.last_name) AS employee_name,
                COALESCE(d.department_name, 'Unassigned') AS department,
                COALESCE(des.designation_name, 'Staff') AS designation,
                COALESCE(pr.final_score, 0.0) AS final_score,
                COALESCE(pr.rating, 'Pending') AS rating,
                COALESCE(pr.recommendation, 'None') AS recommendations
            FROM employees e
            LEFT JOIN departments d ON e.department_id = d.department_id
            LEFT JOIN designations des ON e.designation_id = des.designation_id
            {join_pr_sql}
            {where_sql}
            ORDER BY e.employee_code ASC
        """
        all_params = join_params + params
        cur.execute(query, tuple(all_params))
        rows = cur.fetchall()

        # Calibrate rating & recommendation to strictly match system criteria
        calibrated_rows = []
        for row in rows:
            code, name, dept, desg, score, rating, rec = row
            score_val = float(score or 0.0)
            if score_val > 0:
                if score_val >= 90.0:
                    calibrated_rating = "Outstanding"
                elif score_val >= 80.0:
                    calibrated_rating = "Excellent"
                elif score_val >= 70.0:
                    calibrated_rating = "Good"
                elif score_val >= 60.0:
                    calibrated_rating = "Needs Improvement"
                else:
                    calibrated_rating = "Unsatisfactory"

                if score_val >= 90.0:
                    calibrated_rec = "Promotion Recommended and Salary Increment Recommended"
                elif score_val >= 85.0:
                    calibrated_rec = "Salary Increment Recommended"
                elif score_val >= 75.0:
                    calibrated_rec = "No Immediate Action"
                else:
                    calibrated_rec = "Training / Development Recommended"
            else:
                calibrated_rating = rating or "Pending"
                calibrated_rec = rec or "None"

            calibrated_rows.append((code, name, dept, desg, score_val, calibrated_rating, calibrated_rec))

        return calibrated_rows, selected_cycle_name
    finally:
        cur.close()
        conn.close()


def _log_report_history(user_id, report_name, report_type, file_path):
    """Insert report record into report_history table."""
    conn = get_db_connection()
    cur = conn.cursor()
    report_id = None
    try:
        # Fallback to user_id 1 if not authenticated (e.g. background job/testing)
        uid = user_id if user_id else 1
        cur.execute("""
            INSERT INTO report_history (generated_by, report_name, report_type, generated_on, file_path)
            VALUES (%s, %s, %s, NOW(), %s)
            RETURNING report_id
        """, (uid, report_name, report_type, file_path))
        report_id = cur.fetchone()[0]
        conn.commit()
    except Exception as e:
        conn.rollback()
        print("Error logging report_history:", e)
    finally:
        cur.close()
        conn.close()
    return report_id


@reports_bp.route("/test")
def reports_test():
    return jsonify({
        "success": True,
        "message": "Reports API is working."
    })


@reports_bp.route("/generate-pdf", methods=["GET", "POST"])
def generate_pdf_report():
    """
    Generate professional branded ReportLab PDF and log in report_history.
    """
    title = request.args.get("title") or (request.json.get("title") if request.is_json else None) or "Employee Performance Report"
    title = title.strip()
    cycle_id = request.args.get("cycle_id", type=int) or (request.json.get("cycle_id") if request.is_json else None)

    user_id = session.get("user_id", 1)
    username = session.get("username", "Administrator")
    role_name = session.get("role_name", "HR Admin")
    employee_id = session.get("employee_id")

    timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_title = secure_filename(title.replace(" ", "_"))
    filename = f"{safe_title}_{timestamp_str}.pdf"
    file_path = os.path.join(REPORTS_DIR, filename)

    # 1. Fetch live data
    data_rows, cycle_name = _get_appraisal_report_data(user_id=user_id, role_name=role_name, employee_id=employee_id, cycle_id=cycle_id)

    # 2. Build ReportLab Document
    doc = SimpleDocTemplate(
        file_path,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    muthoot_red = colors.HexColor("#C8102E")
    muthoot_gold = colors.HexColor("#D4AF37")
    text_dark = colors.HexColor("#212529")

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        textColor=muthoot_red,
        spaceAfter=4
    )

    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        textColor=colors.HexColor("#6c757d"),
        spaceAfter=15
    )

    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=11,
        textColor=text_dark
    )

    header_cell_style = ParagraphStyle(
        'HeaderCell',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=11,
        textColor=colors.white
    )

    elements = []

    # Title & Header
    elements.append(Paragraph("MUTHOOT FINANCE - EMPLOYEE APPRAISAL SYSTEM", title_style))
    elements.append(Paragraph(f"Official Report: <b>{title}</b> &bull; Cycle: <b>{cycle_name}</b> &bull; Generated on: {datetime.now().strftime('%d %B %Y, %I:%M %p')}", subtitle_style))

    # Meta Table
    meta_info = [
        [
            Paragraph(f"<b>Requested By:</b> {username} ({role_name})", cell_style),
            Paragraph(f"<b>Total Records:</b> {len(data_rows)}", cell_style),
            Paragraph(f"<b>Export Format:</b> ReportLab PDF", cell_style)
        ]
    ]
    meta_table = Table(meta_info, colWidths=[200, 170, 170])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8F9FA")),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#DEE2E6")),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 10),
        ('RIGHTPADDING', (0, 0), (-1, -1), 10),
    ]))
    elements.append(meta_table)
    elements.append(Spacer(1, 14))

    # Data Table
    table_data = [[
        Paragraph("Emp Code", header_cell_style),
        Paragraph("Employee Name", header_cell_style),
        Paragraph("Department", header_cell_style),
        Paragraph("Designation", header_cell_style),
        Paragraph("Score (%)", header_cell_style),
        Paragraph("Rating", header_cell_style),
        Paragraph("Recommendation", header_cell_style)
    ]]

    for row in data_rows:
        table_data.append([
            Paragraph(str(row[0]), cell_style),
            Paragraph(str(row[1]), cell_style),
            Paragraph(str(row[2]), cell_style),
            Paragraph(str(row[3]), cell_style),
            Paragraph(f"{float(row[4]):.2f}%", cell_style),
            Paragraph(str(row[5]), cell_style),
            Paragraph(str(row[6]), cell_style)
        ])

    data_table = Table(table_data, colWidths=[65, 95, 80, 85, 60, 75, 80])
    data_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), muthoot_red),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9F9F9")]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
    ]))
    elements.append(data_table)

    # Footer note
    elements.append(Spacer(1, 16))
    elements.append(Paragraph("<font color='#888888' size=8>This is a computer-generated appraisal document from the Muthoot Finance Performance Management System. All recorded scores and reviews are confidential.</font>", subtitle_style))

    doc.build(elements)

    # 3. Log to report_history
    rel_path = f"reports/{filename}"
    _log_report_history(user_id, title, "PDF", rel_path)

    return send_file(
        file_path,
        as_attachment=True,
        download_name=filename,
        mimetype="application/pdf"
    )


@reports_bp.route("/generate-excel", methods=["GET", "POST"])
def generate_excel_report():
    """
    Generate professional styled OpenPyXL Excel spreadsheet and log in report_history.
    """
    title = request.args.get("title") or (request.json.get("title") if request.is_json else None) or "Appraisal_Summary_Report"
    title = title.strip()
    cycle_id = request.args.get("cycle_id", type=int) or (request.json.get("cycle_id") if request.is_json else None)

    user_id = session.get("user_id", 1)
    username = session.get("username", "Administrator")
    role_name = session.get("role_name", "HR Admin")
    employee_id = session.get("employee_id")

    timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_title = secure_filename(title.replace(" ", "_"))
    filename = f"{safe_title}_{timestamp_str}.xlsx"
    file_path = os.path.join(REPORTS_DIR, filename)

    # 1. Fetch live data
    data_rows, cycle_name = _get_appraisal_report_data(user_id=user_id, role_name=role_name, employee_id=employee_id, cycle_id=cycle_id)

    # 2. Build OpenPyXL Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Appraisal Summary"

    # Header styling
    muthoot_red_fill = PatternFill(start_color="C8102E", end_color="C8102E", fill_type="solid")
    light_fill = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
    alt_fill = PatternFill(start_color="FAFBFC", end_color="FAFBFC", fill_type="solid")

    title_font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    meta_font = Font(name="Calibri", size=10, italic=True, color="495057")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10, color="212529")

    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    # Row 1: Banner Title
    ws.merge_cells("A1:G1")
    title_cell = ws["A1"]
    title_cell.value = "MUTHOOT FINANCE - EMPLOYEE PERFORMANCE APPRAISAL EXPORT"
    title_cell.fill = muthoot_red_fill
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # Row 2: Metadata Info
    ws.merge_cells("A2:G2")
    meta_cell = ws["A2"]
    meta_cell.value = f"Report: {title} | Cycle: {cycle_name} | Generated On: {datetime.now().strftime('%d-%b-%Y %I:%M %p')} | Exported By: {username} ({role_name})"
    meta_cell.font = meta_font
    meta_cell.fill = light_fill
    meta_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    # Blank Row 3
    ws.row_dimensions[3].height = 10

    # Row 4: Column Headers
    headers = [
        "Employee Code",
        "Employee Name",
        "Department",
        "Designation",
        "Final Score (%)",
        "Appraisal Rating",
        "Promotion Recommendation"
    ]
    for col_idx, header_text in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=header_text)
        cell.fill = muthoot_red_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[4].height = 26

    # Data Rows
    for row_idx, row in enumerate(data_rows, 5):
        current_fill = alt_fill if row_idx % 2 == 0 else PatternFill(fill_type=None)
        values = [
            row[0],
            row[1],
            row[2],
            row[3],
            round(float(row[4]), 2),
            row[5],
            row[6]
        ]
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if current_fill.fill_type:
                cell.fill = current_fill
            if col_idx in [1, 5, 6]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            if col_idx == 5:
                cell.number_format = '0.00'
        ws.row_dimensions[row_idx].height = 20

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 2 and cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(file_path)

    # 3. Log to report_history
    rel_path = f"reports/{filename}"
    _log_report_history(user_id, title, "Excel", rel_path)

    return send_file(
        file_path,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@reports_bp.route("/history", methods=["GET"])
def get_report_history():
    """Retrieve logged generated reports from report_history table."""
    conn = get_db_connection()
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT 
                rh.report_id,
                rh.report_name,
                rh.report_type,
                rh.generated_on,
                rh.file_path,
                COALESCE(u.username, 'System') AS generated_by_user
            FROM report_history rh
            LEFT JOIN users u ON rh.generated_by = u.user_id
            ORDER BY rh.generated_on DESC
            LIMIT 50
        """)
        rows = cur.fetchall()
        reports = [
            {
                "report_id": r[0],
                "report_name": r[1],
                "report_type": r[2],
                "generated_on": r[3].strftime("%d %b %Y, %I:%M %p") if r[3] else None,
                "file_path": r[4],
                "generated_by": r[5]
            }
            for r in rows
        ]
        return jsonify({"success": True, "reports": reports})
    finally:
        cur.close()
        conn.close()


@reports_bp.route("/download/<int:report_id>", methods=["GET"])
def download_saved_report(report_id):
    """Download an existing report from disk using report_id."""
    conn = get_db_connection()
    cur = conn.cursor()
    try:
        cur.execute("SELECT file_path, report_name, report_type FROM report_history WHERE report_id = %s", (report_id,))
        row = cur.fetchone()
        if not row:
            return jsonify({"success": False, "message": "Report record not found."}), 404

        file_path, report_name, report_type = row
        filename = os.path.basename(file_path)
        disk_path = os.path.join(REPORTS_DIR, filename)

        if not os.path.exists(disk_path):
            return jsonify({"success": False, "message": "Report file is no longer available on disk."}), 404

        mime = "application/pdf" if report_type.upper() == "PDF" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        return send_file(disk_path, as_attachment=True, download_name=filename, mimetype=mime)
    finally:
        cur.close()
        conn.close()